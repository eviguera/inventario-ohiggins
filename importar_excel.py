"""
Script para importar productos desde el Excel a productos.json
Procesa todas las hojas relevantes del archivo INVENTARIO ULTIMO.xlsx
"""
import json
import os
import sys

try:
    import openpyxl
except ImportError:
    print("Instalando openpyxl...")
    os.system(f"{sys.executable} -m pip install openpyxl -q")
    import openpyxl

EXCEL_FILE = 'INVENTARIO ULTIMO.xlsx'
OUTPUT_FILE = 'productos.json'

# Hojas que contienen datos de vidrios
HOJAS_VIDRIOS = ['PB  LUNETA', 'NUEVO VP', 'INVENTARIO CYNDY ']

# Mapeo de columnas posibles
MAPEO = {
    'MARCA': ['marca', 'brand', 'marc'],
    'modelo': ['modelo', 'model', 'modelo'],
    'Año': ['año', 'ano', 'year', 'año'],
    'Atril': ['atril', 'racket', 'ubicacion', 'ubicación', 'posicion', 'posición'],
    'VIDRIO': ['vidrio', 'glass', 'cristal', 'tipo'],
    'OBSERVACION': ['observacion', 'observación', 'obs', 'nota', 'note', 'comentario'],
    'Stock': ['stock', 'cantidad', 'qty', 'quantity', 'invent'],
    'COSTO_COMPRA': ['costo compra', 'costo', 'cost', 'precio compra', 'compra', 'costo_compra'],
    'PRECIO_VENTA': ['precio venta', 'precio', 'price', 'venta', 'precio_vent', 'pvp', 'precio vent'],
}


def encontrar_columna(headers, posibles):
    """Busca una columna por posibles nombres"""
    for i, h in enumerate(headers):
        h_lower = str(h or '').lower().strip()
        for p in posibles:
            if h_lower == p or p in h_lower:
                return i
    return -1


def safe_str(v):
    if v is None:
        return ''
    if isinstance(v, (int, float)):
        return str(v).strip()
    return str(v).strip()


def safe_num(v):
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace('$', '').replace(',', '').strip())
    except (ValueError, TypeError):
        return 0


def leer_hoja_vidrios(ws, nombre_hoja):
    """Lee una hoja de vidrios y retorna lista de productos"""
    productos = []
    
    # Encontrar la fila de headers (puede estar en fila 0, 1 o 2)
    header_row = None
    headers = None
    
    for i, row in enumerate(ws.iter_rows(max_row=5, values_only=True)):
        row_str = [str(c or '').lower().strip() for c in row]
        # Buscar fila que tenga al menos 3 de las columnas clave
        matches = sum(1 for key in ['marca', 'modelo', 'atril', 'vidrio', 'stock'] 
                     if any(key in cell for cell in row_str))
        if matches >= 3:
            header_row = i
            headers = [str(c or '').strip() for c in row]
            break
    
    if headers is None:
        print(f"  [WARN] No se encontraron headers en hoja '{nombre_hoja}'")
        return productos
    
    # Mapear columnas
    col_map = {}
    for key, posibles in MAPEO.items():
        col_map[key] = encontrar_columna(headers, posibles)
    
    print(f"  Headers en fila {header_row}: {headers[:9]}")
    print(f"  Columnas mapeadas: {col_map}")
    
    # Leer datos
    for i, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True)):
        if not row or all(c is None or str(c).strip() == '' for c in row):
            continue
        
        # Verificar que tenga al menos marca o modelo
        marca_idx = col_map.get('MARCA', -1)
        modelo_idx = col_map.get('modelo', -1)
        
        marca = safe_str(row[marca_idx]) if marca_idx >= 0 and marca_idx < len(row) else ''
        modelo = safe_str(row[modelo_idx]) if modelo_idx >= 0 and modelo_idx < len(row) else ''
        
        if not marca and not modelo:
            continue
        
        producto = {
            'MARCA': marca.upper(),
            'modelo': modelo,
            'Año': safe_str(row[col_map['Año']]) if col_map.get('Año', -1) >= 0 and col_map['Año'] < len(row) else '',
            'Atril': safe_str(row[col_map['Atril']]) if col_map.get('Atril', -1) >= 0 and col_map['Atril'] < len(row) else '',
            'VIDRIO': safe_str(row[col_map['VIDRIO']]) if col_map.get('VIDRIO', -1) >= 0 and col_map['VIDRIO'] < len(row) else '',
            'OBSERVACION': safe_str(row[col_map['OBSERVACION']]) if col_map.get('OBSERVACION', -1) >= 0 and col_map['OBSERVACION'] < len(row) else '',
            'Stock': int(safe_num(row[col_map['Stock']])) if col_map.get('Stock', -1) >= 0 and col_map['Stock'] < len(row) else 0,
            'COSTO_COMPRA': safe_num(row[col_map['COSTO_COMPRA']]) if col_map.get('COSTO_COMPRA', -1) >= 0 and col_map['COSTO_COMPRA'] < len(row) else 0,
            'PRECIO_VENTA': safe_num(row[col_map['PRECIO_VENTA']]) if col_map.get('PRECIO_VENTA', -1) >= 0 and col_map['PRECIO_VENTA'] < len(row) else 0,
            'HOJA_ORIGEN': nombre_hoja.strip(),
        }
        
        # Limpiar datos
        producto['MARCA'] = producto['MARCA'].replace('  ', ' ').strip()
        producto['modelo'] = producto['modelo'].replace('  ', ' ').strip()
        
        # Filtrar basura
        if producto['MARCA'] in ['', 'MARCA', 'DESCONOCIDO', 'DESCONOCIDOS'] and \
           producto['modelo'] in ['', 'MODELO', 'DESCONOCIDO', 'DESCONOCIDOS']:
            continue
        
        productos.append(producto)
    
    return productos


def main():
    if not os.path.exists(EXCEL_FILE):
        print(f"ERROR: No se encontró {EXCEL_FILE}")
        return
    
    print(f"Leyendo {EXCEL_FILE}...")
    wb = openpyxl.load_workbook(EXCEL_FILE, read_only=True, data_only=True)
    
    todos_los_productos = []
    
    for hoja in HOJAS_VIDRIOS:
        if hoja not in wb.sheetnames:
            print(f"[WARN] Hoja '{hoja}' no encontrada")
            continue
        
        print(f"\nProcesando hoja: {hoja}")
        ws = wb[hoja]
        productos = leer_hoja_vidrios(ws, hoja)
        print(f"  [OK] {len(productos)} productos encontrados")
        todos_los_productos.extend(productos)
    
    # Eliminar duplicados (misma marca, modelo, año, atril, vidrio)
    print(f"\nTotal antes de eliminar duplicados: {len(todos_los_productos)}")
    vistos = set()
    unicos = []
    for p in todos_los_productos:
        key = (p['MARCA'], p['modelo'], p['Año'], p['Atril'], p['VIDRIO'])
        if key not in vistos:
            vistos.add(key)
            unicos.append(p)
    
    print(f"Total después de eliminar duplicados: {len(unicos)}")
    
    # Guardar
    with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
        json.dump(unicos, f, ensure_ascii=False, indent=2)
    
    print(f"\n[OK] {len(unicos)} productos guardados en {OUTPUT_FILE}")
    
    # Estadísticas
    stock_total = sum(p['Stock'] for p in unicos)
    con_stock = sum(1 for p in unicos if p['Stock'] > 0)
    sin_stock = sum(1 for p in unicos if p['Stock'] == 0)
    valor_total = sum(p['COSTO_COMPRA'] * p['Stock'] for p in unicos)
    
    print(f"\n--- ESTADÍSTICAS ---")
    print(f"Productos con stock: {con_stock}")
    print(f"Productos sin stock: {sin_stock}")
    print(f"Unidades totales: {stock_total}")
    print(f"Valor inventario (costo): ${valor_total:,.0f}")


if __name__ == '__main__':
    main()
